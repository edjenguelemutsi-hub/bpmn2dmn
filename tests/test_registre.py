# -*- coding: utf-8 -*-
"""
Tests du registre des expérimentations : les mesures d'une extraction, le
code de référence du modèle, l'écriture CSV et Excel, la conservation des
sorties, la commande `registre`, et l'enregistrement automatique par
l'application web.
"""
import csv
import json
import os
import shutil
import sys
import tempfile
import threading
import unittest
import urllib.request
import zipfile
import xml.etree.ElementTree as ET

ICI = os.path.dirname(os.path.abspath(__file__))
RACINE = os.path.dirname(ICI)
sys.path.insert(0, RACINE)

from bpmn2dmn.parseur import parser_bpmn
from bpmn2dmn.transformation import transformer
from bpmn2dmn.serialiseur import serialiser_dmn
from bpmn2dmn.qualification import Arbitrages
from bpmn2dmn.metriques import COLONNES, NOMS_COLONNES, mesurer
from bpmn2dmn.registre import (Registre, abreger, code_modele, empreinte,
                               texte_synthese)
from bpmn2dmn.classeur import ecrire_classeur

FIMEX = os.path.join(RACINE, "corpus", "fimex.bpmn")
COPD = os.path.join(RACINE, "corpus", "copd.bpmn")
ARBITRAGES_FIMEX = os.path.join(RACINE, "corpus", "fimex.arbitrages.json")


class TestMesures(unittest.TestCase):
    """Les grandeurs mesurées sur FIMEX sont celles des travaux antérieurs."""

    @classmethod
    def setUpClass(cls):
        cls.mi = parser_bpmn(FIMEX)
        cls.mf = transformer(cls.mi)
        cls.dmn = serialiser_dmn(cls.mf)
        cls.ligne = mesurer(cls.mi, cls.mf, dmn=cls.dmn,
                            reference=["Affecter le dossier à un agent",
                                       "Vérifier la conformité documentaire",
                                       "Consulter la demande et la décision"])

    def test_colonnes_uniques_et_ordonnees(self):
        self.assertEqual(len(NOMS_COLONNES), len(set(NOMS_COLONNES)))
        self.assertEqual(list(self.ligne.keys()), NOMS_COLONNES)
        for nom, groupe, libelle, _ in COLONNES:
            self.assertTrue(nom and groupe and libelle, nom)

    def test_fimex_sans_arbitrage(self):
        l = self.ligne
        self.assertEqual(l["decisions"], 3)
        self.assertEqual((l["ir"], l["kr"], l["ar"]), (5, 1, 1))
        self.assertEqual(l["points_expert"], 10)
        self.assertEqual(l["da_candidats"], 9)          # |DA| = 9, |DA*| = 3
        self.assertEqual(l["da_retenus"], 3)
        self.assertEqual(l["pe_candidats"], 6)
        self.assertEqual(l["pe_annotations"], 2)
        self.assertEqual(l["pe_ressources"], 2)
        self.assertEqual(l["acyclique"], 1)
        self.assertEqual(l["invariants_violes"], 0)
        self.assertEqual(l["dmn_relu_conforme"], 1)
        self.assertEqual(l["precision"], 1.0)
        self.assertEqual(l["rappel"], 1.0)
        self.assertEqual(l["questions_par_decision"], round(10 / 3, 2))
        self.assertEqual(l["vf_sommets"], l["decisions"] + l["donnees_entree"]
                         + l["sources_connaissance"] + l["bkm"])
        self.assertEqual(l["ef_exigences"], l["ir"] + l["kr"] + l["ar"])
        self.assertEqual(l["ir"], l["ir_donnee_decision"] + l["ir_decision_decision"])

    def test_fimex_avec_arbitrages(self):
        arb = Arbitrages.depuis_json(ARBITRAGES_FIMEX)
        mf = transformer(self.mi, arb)
        l = mesurer(self.mi, mf, arbitrages=arb,
                    reference=["Affecter le dossier à un agent",
                               "Vérifier la conformité documentaire",
                               "Consulter la demande et la décision"])
        self.assertEqual((l["vf_sommets"], l["ef_exigences"]), (10, 8))
        self.assertEqual((l["ir"], l["kr"], l["ar"]), (5, 1, 2))
        self.assertEqual(l["p4_dependances"], 1)
        # la référence nomme les activités : un renommage n'est pas une erreur
        self.assertEqual(l["precision"], 1.0)
        self.assertEqual(l["rappel"], 1.0)

    def test_structure_du_drd(self):
        l = self.ligne
        self.assertGreaterEqual(l["drd_profondeur"], 1)
        self.assertGreaterEqual(l["drd_composantes"], 1)
        self.assertGreater(l["drd_densite"], 0)
        self.assertGreaterEqual(l["drd_degre_entrant_max"], 1)

    def test_bpco(self):
        mi = parser_bpmn(COPD)
        mf = transformer(mi)
        l = mesurer(mi, mf)
        self.assertEqual(l["decisions"], 4)
        self.assertEqual(l["points_expert"], 9)
        self.assertEqual(l["p3_optionnels"], 1)          # « Free time slot »


class TestCodeModele(unittest.TestCase):
    def test_stable_et_sensible_au_contenu(self):
        self.assertEqual(code_modele("Corpus/FIMEX v2.bpmn", "<a/>"),
                         code_modele("fimex-v2.bpmn", "<a/>\r\n"[:-2]))
        self.assertEqual(abreger("Corpus/FIMEX v2.bpmn"), "fimex-v2")
        self.assertNotEqual(empreinte("<a/>"), empreinte("<b/>"))
        self.assertEqual(empreinte("x\r\ny"), empreinte("x\ny"))
        self.assertEqual(len(empreinte("")), 8)
        self.assertEqual(abreger(""), "modele")


class TestRegistre(unittest.TestCase):
    def setUp(self):
        self.dossier = tempfile.mkdtemp(prefix="registre-")
        self.registre = Registre(self.dossier)

    def tearDown(self):
        shutil.rmtree(self.dossier, ignore_errors=True)

    def _enregistrer(self, chemin=FIMEX, contexte="test", **kw):
        mi = parser_bpmn(chemin)
        mf = transformer(mi)
        return self.registre.enregistrer(mi, mf, contexte, fichier=chemin,
                                         dmn=serialiser_dmn(mf),
                                         rapport={"resume": mf.resume()}, **kw)

    def test_une_ligne_par_extraction(self):
        l1 = self._enregistrer()
        l2 = self._enregistrer(COPD, contexte="validation corpus")
        self.assertEqual(l1["experience"], "EXP-000001")
        self.assertEqual(l2["experience"], "EXP-000002")
        self.assertTrue(l1["code_modele"].startswith("fimex-"))
        self.assertEqual(l1["precision"], 1.0)          # .gold voisin lu
        lignes = self.registre.lignes()
        self.assertEqual(len(lignes), 2)
        self.assertEqual(lignes[0]["decisions"], 3)
        self.assertEqual(lignes[1]["decisions"], 4)
        self.assertEqual(lignes[1]["contexte"], "validation corpus")
        with open(self.registre.chemin_csv, encoding="utf-8-sig") as f:
            entete = next(csv.reader(f, delimiter=";"))
        self.assertEqual(entete, NOMS_COLONNES)

    def test_sorties_conservees(self):
        l = self._enregistrer()
        base = os.path.join(self.registre.dossier_sorties,
                            "%s_%s" % (l["experience"], l["code_modele"]))
        for suffixe in (".bpmn", ".drd.dmn", ".rapport.json"):
            self.assertTrue(os.path.isfile(base + suffixe), suffixe)
        with open(base + ".rapport.json", encoding="utf-8") as f:
            self.assertEqual(json.load(f)["resume"]["decisions"], 3)
        self.assertTrue(os.path.isfile(os.path.join(self.dossier,
                                                    "dictionnaire-colonnes.md")))

    def test_classeur_excel(self):
        self._enregistrer()
        self._enregistrer(COPD)
        chemin = self.registre.chemin_xlsx
        self.assertTrue(os.path.isfile(chemin))
        with zipfile.ZipFile(chemin) as z:
            noms = z.namelist()
            self.assertIn("xl/workbook.xml", noms)
            self.assertIn("xl/worksheets/sheet1.xml", noms)
            self.assertIn("xl/worksheets/sheet3.xml", noms)
            feuille = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
            ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            lignes = feuille.findall(".//m:sheetData/m:row", ns)
            self.assertEqual(len(lignes), 3)             # en-tête + 2
            cellules = lignes[0].findall("m:c", ns)
            self.assertEqual(len(cellules), len(NOMS_COLONNES))
            for f in z.namelist():
                if f.endswith(".xml") or f.endswith(".rels"):
                    ET.fromstring(z.read(f))             # XML bien formé
        try:
            import openpyxl
        except ImportError:
            return
        wb = openpyxl.load_workbook(chemin)
        self.assertEqual(wb.sheetnames[:3], ["Expérimentations", "Synthèse",
                                             "Dictionnaire"])
        ws = wb["Expérimentations"]
        self.assertEqual(ws["A1"].value, "experience")
        self.assertEqual(ws["A2"].value, "EXP-000001")
        colonne_decisions = NOMS_COLONNES.index("decisions") + 1
        self.assertEqual(ws.cell(row=2, column=colonne_decisions).value, 3)

    def test_synthese(self):
        self._enregistrer()
        self._enregistrer(COPD)
        s = self.registre.synthese()
        self.assertEqual(s["nombre"], 2)
        self.assertEqual(s["modeles_distincts"], 2)
        self.assertEqual(s["moyennes"]["decisions"], 3.5)
        texte = texte_synthese(self.registre)
        self.assertIn("expérimentations enregistrées ... 2", texte)

    def test_campagne(self):
        self.registre.journal_campagne("campagne", 1, 10, 10, 30, 70, 40, 0, 1.2)
        self.assertEqual(len(self.registre.campagnes()), 1)
        self.assertEqual(self.registre.campagnes()[0]["decisions"], 30)
        with zipfile.ZipFile(self.registre.chemin_xlsx) as z:
            self.assertIn("xl/worksheets/sheet4.xml", z.namelist())

    def test_ancien_format_conserve(self):
        os.makedirs(self.dossier, exist_ok=True)
        with open(self.registre.chemin_csv, "w", encoding="utf-8-sig") as f:
            f.write("experience;autre\nEXP-000001;x\n")
        l = self._enregistrer()
        # la numérotation continue : un numéro n'est jamais réattribué
        self.assertEqual(l["experience"], "EXP-000002")
        anciens = [n for n in os.listdir(self.dossier)
                   if n.startswith("registre-ancien-")]
        self.assertEqual(len(anciens), 1)

    def test_ecritures_concurrentes(self):
        erreurs = []

        def travail():
            try:
                self._enregistrer()
            except Exception as e:
                erreurs.append(e)
        fils = [threading.Thread(target=travail) for _ in range(6)]
        for f in fils:
            f.start()
        for f in fils:
            f.join()
        self.assertEqual(erreurs, [])
        numeros = [l["experience"] for l in self.registre.lignes()]
        self.assertEqual(numeros, ["EXP-%06d" % i for i in range(1, 7)])

    def test_classeur_minimal(self):
        chemin = os.path.join(self.dossier, "essai.xlsx")
        ecrire_classeur(chemin, [("Feuille", ("a", "b"), [[1, 2.5], ["x", ""]], None)])
        with zipfile.ZipFile(chemin) as z:
            self.assertIn("[Content_Types].xml", z.namelist())


class TestCommandeRegistre(unittest.TestCase):
    def setUp(self):
        self.dossier = tempfile.mkdtemp(prefix="registre-cli-")

    def tearDown(self):
        shutil.rmtree(self.dossier, ignore_errors=True)

    def test_transformer_puis_registre(self):
        import cli
        import io
        from contextlib import redirect_stdout
        sortie = os.path.join(self.dossier, "fimex.drd.dmn")
        with redirect_stdout(io.StringIO()) as tampon:
            code = cli.principal(["transformer", FIMEX, "-o", sortie,
                                  "--registre", self.dossier])
        self.assertEqual(code, 0)
        self.assertIn("EXP-000001", tampon.getvalue())
        with redirect_stdout(io.StringIO()) as tampon:
            code = cli.principal(["registre", "--registre", self.dossier,
                                  "--dernieres", "3"])
        self.assertEqual(code, 0)
        self.assertIn("expérimentations enregistrées ... 1", tampon.getvalue())
        with redirect_stdout(io.StringIO()) as tampon:
            cli.principal(["transformer", FIMEX, "-o", sortie,
                           "--registre", self.dossier, "--sans-registre"])
        self.assertNotIn("EXP-000002", tampon.getvalue())
        self.assertEqual(len(Registre(self.dossier).lignes()), 1)


class TestApplicationWeb(unittest.TestCase):
    """L'application web consigne chaque extraction et sert le registre."""

    @classmethod
    def setUpClass(cls):
        import app
        cls.dossier = tempfile.mkdtemp(prefix="registre-web-")
        cls.ancien = app.REGISTRE
        app.REGISTRE = Registre(cls.dossier)
        cls.serveur = app.Serveur(("127.0.0.1", 0), app.Requetes)
        cls.port = cls.serveur.server_address[1]
        cls.fil = threading.Thread(target=cls.serveur.serve_forever, daemon=True)
        cls.fil.start()
        cls.ouvreur = urllib.request.build_opener(urllib.request.ProxyHandler({}))

    @classmethod
    def tearDownClass(cls):
        import app
        cls.serveur.shutdown()
        cls.serveur.server_close()
        app.REGISTRE = cls.ancien
        shutil.rmtree(cls.dossier, ignore_errors=True)

    def _post(self, chemin, corps):
        req = urllib.request.Request(
            "http://127.0.0.1:%d%s" % (self.port, chemin),
            data=json.dumps(corps).encode("utf-8"),
            headers={"Content-Type": "application/json"})
        with self.ouvreur.open(req, timeout=30) as r:
            return json.loads(r.read().decode("utf-8"))

    def _get(self, chemin):
        with self.ouvreur.open("http://127.0.0.1:%d%s" % (self.port, chemin),
                               timeout=30) as r:
            return r.status, r.headers, r.read()

    def test_extraction_consignee_et_registre_servi(self):
        with open(FIMEX, encoding="utf-8") as f:
            texte = f.read()
        d = self._post("/api/transformer", {"bpmn": texte, "nom": "fimex.bpmn"})
        self.assertTrue(d["ok"])
        x = d["rapport"]["experimentation"]
        self.assertEqual(x["experience"], "EXP-000001")
        self.assertTrue(x["code_modele"].startswith("fimex-"))
        statut, _, corps = self._get("/api/registre")
        r = json.loads(corps.decode("utf-8"))
        self.assertTrue(r["actif"])
        self.assertEqual(r["nombre"], 1)
        self.assertEqual(r["dernieres"][0]["fichier"], "fimex.bpmn")
        self.assertEqual(r["dernieres"][0]["contexte"], "application web")
        statut, entetes, corps = self._get("/registre/registre.csv")
        self.assertEqual(statut, 200)
        self.assertIn("attachment", entetes["Content-Disposition"])
        self.assertTrue(corps.decode("utf-8-sig").startswith("experience;"))
        statut, entetes, corps = self._get("/registre/registre.xlsx")
        self.assertEqual(statut, 200)
        self.assertEqual(corps[:2], b"PK")
        with self.assertRaises(urllib.error.HTTPError):
            self._get("/registre/../app.py")

    def test_page_presente_le_registre(self):
        statut, _, corps = self._get("/")
        page = corps.decode("utf-8")
        self.assertIn("Registre des expérimentations", page)
        self.assertIn("/registre/registre.xlsx", page)


if __name__ == "__main__":
    unittest.main()
